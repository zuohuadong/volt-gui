#!/usr/bin/env python3
"""本地生成失效分析报告草稿 DOCX（确定性排版，不联网）。

输入是一个 JSON 规格文件，由 anyong-integrations-mcp.mjs 的 report_generate_docx 工具生成。
只负责排版：标题、草稿声明、项目信息、六节正文、本地证据清单、缺口和来源行。
所有内容都来自调用方提供的规格，脚本不做推断、不补造数据。
"""

from __future__ import annotations

import hashlib
import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm


def fail(message: str) -> None:
    print(json.dumps({"error": message}, ensure_ascii=False), flush=True)
    raise SystemExit(1)


def _human_size(value) -> str:
    try:
        size = float(value)
    except (TypeError, ValueError):
        return ""
    if size < 1024:
        return f"{int(size)} B"
    for unit in ("KB", "MB", "GB"):
        size /= 1024.0
        if size < 1024 or unit == "GB":
            return f"{size:.1f} {unit}"
    return ""


def set_font(run, name: str = "宋体", size: float = 10.5, bold: bool = False, color=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def heading(document, text: str, level: int) -> None:
    paragraph = document.add_heading(level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_font(run, "黑体", {1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True,
             color=RGBColor(0x1F, 0x2A, 0x37))


def body(document, text: str) -> None:
    for block in [item for item in str(text).split("\n") if item.strip()]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(21)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.4
        set_font(paragraph.add_run(block.strip()))


def table(document, headers, rows) -> None:
    grid = document.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = grid.rows[0].cells[index]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(str(header)), "黑体", 10, bold=True)
    for row in rows:
        cells = grid.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            set_font(cells[index].paragraphs[0].add_run(str(value)), "宋体", 9.5)


def main() -> None:
    if len(sys.argv) != 2:
        fail("用法：fa-report-docx.py <spec.json>")
    spec_path = sys.argv[1]
    try:
        with open(spec_path, "r", encoding="utf-8") as handle:
            spec = json.load(handle)
    except (OSError, json.JSONDecodeError) as error:
        fail(f"规格文件无法读取：{error}")

    output = str(spec.get("output") or "").strip()
    if not output or not os.path.isabs(output):
        fail("output 必须是绝对路径")
    if os.path.exists(output) and spec.get("overwrite") is not True:
        fail("目标文件已存在；如需覆盖请显式传 overwrite=true")
    os.makedirs(os.path.dirname(output), exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.2))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))

    title = str(spec.get("title") or "失效分析报告（草稿）").strip()
    heading(document, title, 1)
    notice = str(spec.get("notice") or "").strip()
    if notice:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.add_run(notice), "黑体", 10.5, bold=True,
                 color=RGBColor(0xB0, 0x1C, 0x1C))

    identifiers = spec.get("identifiers") or []
    if identifiers:
        heading(document, "项目信息", 2)
        table(document, ["项目", "值"], [(item.get("label", ""), item.get("value", "")) for item in identifiers])

    for item in spec.get("sections") or []:
        heading(document, str(item.get("title") or item.get("key") or "章节"), 2)
        body(document, str(item.get("body") or ""))

    evidence = spec.get("evidence") or []
    if evidence:
        heading(document, "本地证据清单", 2)
        table(
            document,
            ["序号", "文件", "类型", "大小", "修改时间", "相对路径"],
            [
                (
                    str(index + 1),
                    item.get("fileName", "") or item.get("title", ""),
                    item.get("evidenceType", ""),
                    _human_size(item.get("sizeBytes")),
                    str(item.get("modifiedAt", ""))[:19].replace("T", " "),
                    item.get("relativePath", ""),
                )
                for index, item in enumerate(evidence)
            ],
        )
        paragraph = document.add_paragraph()
        set_font(
            paragraph.add_run("以上文件由本机资料目录扫描得到，未逐份核对内容；引用前需人工确认。"),
            "宋体",
            9,
            color=RGBColor(0x55, 0x5F, 0x6D),
        )

    gaps = [str(item).strip() for item in (spec.get("gaps") or []) if str(item).strip()]
    heading(document, "资料缺口与待确认", 2)
    if gaps:
        for item in gaps:
            paragraph = document.add_paragraph(style="List Bullet")
            set_font(paragraph.add_run(item))
    else:
        body(document, "本次未收到 FA 证据预检缺口；仍以 FA 审核结论为准。")

    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(
        footer.add_run(str(spec.get("source") or "")),
        "宋体",
        9,
        color=RGBColor(0x55, 0x5F, 0x6D),
    )

    document.core_properties.title = title
    document.core_properties.comments = str(spec.get("source") or "")
    document.save(output)

    with open(output, "rb") as handle:
        payload = handle.read()
    print(json.dumps({
        "output": output,
        "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(),
        "sections": len(spec.get("sections") or []),
        "evidence": len(evidence),
        "gaps": len(gaps),
    }, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()