#!/usr/bin/env python3
"""Bounded spreadsheet engine for the Anyong integrations MCP server.

The engine is deliberately small: it only reads spreadsheets and computes
read-only SQL aggregates.  It never writes to the source workbook.

Invocation:
    python sheet-engine.py <request.json>

The request file carries a JSON object; the result is one JSON line on stdout.
Only openpyxl (streaming, memory bounded) and duckdb (SQL) are required, so the
engine runs on the offline Python runtime that ships with the desktop app.
"""

import csv
import datetime
import json
import math
import os
import sys
import tempfile
import time
import traceback
from decimal import Decimal
from pathlib import Path

MAX_PROFILE_ROWS = 200
MAX_PROFILE_COLUMNS = 80
MAX_SQL_ROWS = 5000
MAX_CELL_CHARS = 4000


class SheetError(Exception):
    """A user-facing failure; the message is returned to the model verbatim."""


def read_request() -> dict:
    if len(sys.argv) < 2:
        raise SheetError("缺少请求文件参数")
    return json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))


def open_workbook(path: str):
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - runtime provisioning issue
        raise SheetError(f"Python 运行时缺少 openpyxl：{error}") from error
    if not Path(path).is_file():
        raise SheetError(f"表格文件不存在：{path}")
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise SheetError(f"无法打开表格文件（可能已损坏、加密或不是 xlsx）：{error}") from error


def select_sheet(workbook, requested):
    names = list(workbook.sheetnames)
    if not names:
        raise SheetError("工作簿里没有工作表")
    if not requested:
        return workbook[names[0]], names
    if requested in names:
        return workbook[requested], names
    raise SheetError(f"工作表 {requested} 不存在；可用工作表：{', '.join(names)}")


def sheet_dimensions(sheet):
    """Return (rows, columns); read-only worksheets only expose them lazily."""
    rows = getattr(sheet, "max_row", None)
    columns = getattr(sheet, "max_column", None)
    if rows and columns:
        return rows, columns
    try:
        dimension = sheet.calculate_dimension()
    except Exception:
        return rows, columns
    if not dimension or ":" not in dimension:
        return rows, columns
    start, end = dimension.split(":", 1)
    parsed_start = parse_range(start, 1)
    parsed_end = parse_range(end.replace("$", ""), 1)
    rows = rows or parsed_end[0] + 1
    columns = columns or parsed_end[3] + 1
    return rows, columns


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    text = value if isinstance(value, str) else str(value)
    if len(text) > MAX_CELL_CHARS:
        return text[:MAX_CELL_CHARS]
    return text


def json_value(value):
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, Decimal):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (bytes, bytearray)):
        return f"<{len(value)} bytes>"
    if isinstance(value, (list, tuple)):
        return [json_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_value(item) for key, item in value.items()}
    return str(value)[:MAX_CELL_CHARS]


def column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def parse_range(raw, max_rows: int):
    """Parse an A1:B10 style range into zero-based row/column bounds."""
    if not raw:
        return 0, None, 0, None
    text = str(raw).strip().upper().replace("$", "")
    if "!" in text:
        _, text = text.rsplit("!", 1)
    if ":" not in text:
        text = f"{text}:{text}"

    def split_cell(cell: str):
        digits = "".join(char for char in cell if char.isdigit())
        letters = "".join(char for char in cell if char.isalpha())
        if not digits or not letters:
            raise SheetError(f"无效的单元格范围：{raw}")
        row = int(digits) - 1
        column = 0
        for char in letters:
            column = column * 26 + (ord(char) - 64)
        return row, column - 1

    start_row, start_column = split_cell(text.split(":", 1)[0])
    end_row, end_column = split_cell(text.split(":", 1)[1])
    if end_row < start_row or end_column < start_column:
        raise SheetError(f"范围结束位置早于开始位置：{raw}")
    if end_row - start_row + 1 > max_rows:
        end_row = start_row + max_rows - 1
    return start_row, end_row, start_column, end_column


def describe_columns(rows, limit_columns=MAX_PROFILE_COLUMNS):
    width = max((len(row) for row in rows), default=0)
    header = rows[0] if rows else []
    body = [row for row in rows[1:] if any(cell not in (None, "") for cell in row)]
    columns = []
    for index in range(min(width, limit_columns)):
        name = header[index] if index < len(header) else None
        values = [row[index] if index < len(row) else None for row in body]
        present = [value for value in values if value not in (None, "")]
        kinds = {type(value).__name__ for value in present}
        columns.append({
            "column": column_letter(index),
            "header": cell_text(name) or None,
            "kinds": sorted(kinds)[:3],
            "nulls": len(values) - len(present),
            "distinctSampled": len({cell_text(value) for value in present}) if len(present) <= 5000 else None,
            "sample": [json_value(value) for value in present[:3]],
        })
    return columns, len(body)


def sample_via_calamine(path: str, sheet_name, max_rows: int):
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_path(path)
    names = list(workbook.sheet_names)
    if sheet_name:
        if sheet_name not in names:
            raise SheetError(f"工作表 {sheet_name} 不存在；可用工作表：{', '.join(names)}")
        sheet = workbook.get_sheet_by_name(sheet_name)
        title = sheet_name
    else:
        sheet = workbook.get_sheet_by_index(0)
        title = names[0] if names else ""
    rows = []
    for row in sheet.iter_rows():
        rows.append([json_value(value) for value in row[:MAX_PROFILE_COLUMNS]])
        if len(rows) > max_rows:
            break
    return rows, title, names


def sample_via_duckdb(request: dict, path: str, sheet_name, max_rows: int):
    """Sample a large workbook through the streaming extension (no full parse)."""
    import duckdb

    connection = duckdb.connect(database=":memory:")
    try:
        connection.execute("SET memory_limit='512MB'")
        if not load_excel_extension(connection, request):
            raise SheetError(oversize_message(estimate_cells(path), int(request.get("cellBudget") or DEFAULT_CELL_BUDGET), request))
        selected = sheet_name or ""
        statement = (
            "SELECT * FROM read_xlsx("
            f"{quote_literal(path)}, sheet={quote_literal(str(selected))}, header=true) "
            f"LIMIT {max(1, max_rows)}"
        )
        cursor = connection.execute(statement)
        column_names = [item[0] for item in cursor.description]
        rows = [[json_value(value) for value in row] for row in cursor.fetchall()]
        return [column_names] + rows, selected or (fast_sheet_names(path) or [""])[0], fast_sheet_names(path)
    finally:
        connection.close()


def fast_sheet_names(path: str):
    """Read sheet names from workbook.xml without opening the workbook engine."""
    try:
        import re
        import zipfile

        with zipfile.ZipFile(path) as archive:
            xml = archive.read("xl/workbook.xml").decode("utf-8", "replace")
        return [name for name in re.findall(r'<sheet[^>]*name="([^"]+)"', xml)]
    except Exception:
        return None


def query_from_cache(request: dict, parquet_path: Path, limit: int):
    """Serve an A1-anchored window from the cached snapshot instead of openpyxl."""
    raw = str(request.get("range") or "").strip()
    end_column = None
    if raw:
        start_row, end_row, start_column, end_column = parse_range(raw, limit)
        if start_row != 0 or start_column != 0:
            return None
    else:
        end_row = None
    try:
        import duckdb
    except ImportError:
        return None
    try:
        connection = duckdb.connect(database=":memory:")
    except Exception:
        return None
    try:
        connection.execute("SET memory_limit='256MB'")
        connection.execute(
            f"CREATE VIEW cached AS SELECT * FROM read_parquet({quote_literal(str(parquet_path))})"
        )
        column_names = [row[0] for row in connection.execute("DESCRIBE cached").fetchall()]
        total_width = len(column_names)
        width = total_width if end_column is None else min(end_column + 1, total_width)
        if end_row is None:
            end_row = limit
        wanted = max(0, min(end_row, limit))
        fetched = connection.execute(f"SELECT * FROM cached LIMIT {wanted}").fetchall()
        rows = [{"row": 1, "cells": [str(name) for name in column_names[:width]]}]
        for index, row in enumerate(fetched, start=2):
            rows.append({"row": index, "cells": [json_value(value) for value in row[:width]]})
        return {
            "sheet": request.get("sheet") or "",
            "sheets": fast_sheet_names(str(request["path"])),
            "range": raw or f"A1:{column_letter(width - 1)}{wanted + 1}",
            "columns": [column_letter(index) for index in range(width)],
            "rows": rows,
            "rowCount": len(rows),
            "truncated": len(fetched) >= wanted,
            "source": "cache",
        }
    except Exception:
        return None
    finally:
        connection.close()


def profile(request: dict) -> dict:
    path = str(request["path"])
    max_rows = int(request.get("maxRows") or MAX_PROFILE_ROWS)
    max_rows = max(1, min(max_rows, MAX_PROFILE_ROWS))
    names = fast_sheet_names(path)
    estimate = estimate_cells(path)
    resolved = str(request.get("sheet") or "")
    try:
        # Sampling is cheapest through the streaming extension (LIMIT pushdown),
        # so prefer it whenever it is provisioned, whatever the sheet size.
        if not str(request.get("reader") or "").strip() and duckdb_extension_path(request) is not None:
            reader = "duckdb-excel"
        else:
            reader = resolve_reader(request, path)
    except SheetError as error:
        # Discovery must stay instant even when a full conversion is refused.
        return {
            "sheet": resolved or (names[0] if names else ""),
            "sheets": names,
            "rows": None,
            "columns": None,
            "cellsEstimate": estimate,
            "sampledRows": 0,
            "reader": "blocked",
            "dataColumns": [],
            "warnings": [str(error)],
        }
    if reader == "duckdb-excel":
        rows, title, sheet_names = sample_via_duckdb(request, path, resolved, max_rows + 1)
        columns, sampled = describe_columns(rows)
        return {
            "sheet": title,
            "sheets": sheet_names or names,
            "rows": None,
            "columns": None,
            "cellsEstimate": estimate,
            "sampledRows": sampled,
            "reader": reader,
            "dataColumns": columns,
            "warnings": [f"大表走流式扩展，仅采样前 {sampled} 行；整表约 {estimate:,} 个单元格，请用 sheet_aggregate 统计。" if estimate else "仅采样前若干行"],
        }
    declared_rows, declared_columns = (None, None)
    if reader == "calamine":
        rows, title, sheet_names = sample_via_calamine(path, resolved, max_rows)
    else:
        workbook = open_workbook(path)
        sheet, sheet_names = select_sheet(workbook, resolved)
        title = sheet.title
        declared_rows, declared_columns = sheet_dimensions(sheet)
        rows = [[json_value(value) for value in row[:MAX_PROFILE_COLUMNS]] for row in sheet.iter_rows(max_row=max_rows, values_only=True)]
    columns, sampled = describe_columns(rows)
    warnings = []
    if declared_rows and declared_rows > len(rows):
        warnings.append(f"仅采样前 {len(rows)} 行（工作表共约 {declared_rows} 行）")
    if rows and not any(column["sample"] for column in columns):
        warnings.append("采样区间为空：可能数据不在前几行，请用 sheet_query 指定范围；公式单元格若没有缓存值也会显示为空")
    return {
        "sheet": title,
        "sheets": sheet_names,
        "rows": declared_rows,
        "columns": declared_columns,
        "cellsEstimate": estimate,
        "sampledRows": sampled,
        "reader": reader,
        "dataColumns": columns,
        "warnings": warnings,
    }


def query(request: dict) -> dict:
    limit = int(request.get("limit") or request.get("maxRows") or MAX_SQL_ROWS)
    limit = max(1, min(limit, MAX_SQL_ROWS))
    source = Path(request["path"])
    cache_dir = Path(request.get("cacheDir") or tempfile.gettempdir())
    cached_snapshot = cache_paths(source, str(request.get("sheet") or ""), cache_dir)[1]
    if cached_snapshot.is_file():
        served = query_from_cache(request, cached_snapshot, limit)
        if served is not None:
            return served
    workbook = open_workbook(str(source))
    sheet, names = select_sheet(workbook, request.get("sheet"))
    start_row, end_row, start_column, end_column = parse_range(request.get("range"), limit)
    if end_row is None:
        end_row = start_row + limit - 1
    if end_column is None:
        _, declared_columns = sheet_dimensions(sheet)
        end_column = max((declared_columns or 1) - 1, 0)
    rows = []
    for offset, row in enumerate(sheet.iter_rows(values_only=True)):
        if offset < start_row:
            continue
        if offset > end_row:
            break
        cells = row[start_column:end_column + 1]
        truncated_cells = [json_value(value) for value in cells]
        if any(cell not in (None, "") for cell in truncated_cells):
            rows.append({"row": offset + 1, "cells": truncated_cells})
    columns = [column_letter(index) for index in range(start_column, end_column + 1)]
    return {
        "sheet": sheet.title,
        "sheets": names,
        "range": request.get("range") or f"{column_letter(start_column)}{start_row + 1}:{column_letter(end_column)}{end_row + 1}",
        "columns": columns,
        "rows": rows,
        "rowCount": len(rows),
        "truncated": len(rows) >= limit,
        "source": "openpyxl",
    }


DEFAULT_CELL_BUDGET = 10_000_000
# Rough bytes of worksheet XML per cell; used to bound the Rust reader's memory
# before it materialises anything.
XML_BYTES_PER_CELL = 50


def calamine_available() -> bool:
    try:
        import python_calamine  # noqa: F401
    except Exception:
        return False
    return True


def estimate_cells(path: str):
    """Cheap cell estimate from the zip central directory (no parsing)."""
    try:
        import zipfile

        with zipfile.ZipFile(path) as archive:
            total = sum(
                entry.file_size
                for entry in archive.infolist()
                if entry.filename.startswith("xl/worksheets/")
            )
        if total:
            return int(total / XML_BYTES_PER_CELL)
    except Exception:
        pass
    try:
        return int(Path(path).stat().st_size / 20)
    except Exception:
        return None


def duckdb_platform() -> str:
    import platform

    system = platform.system().lower()
    machine = platform.machine().lower()
    if system == "windows":
        return "windows_amd64" if machine in ("amd64", "x86_64") else "windows_arm64"
    if system == "darwin":
        return "osx_arm64" if machine in ("arm64", "aarch64") else "osx_amd64"
    if system == "linux":
        return "linux_arm64" if machine in ("aarch64", "arm64") else "linux_amd64"
    return ""


def duckdb_extension_path(request: dict):
    """Locate the pinned excel extension without letting duckdb reach the network."""
    explicit = str(request.get("extensionPath") or "").strip()
    if explicit and Path(explicit).is_file():
        return Path(explicit)
    directory = str(request.get("extensionDirectory") or "").strip()
    if not directory:
        return None
    try:
        import duckdb

        version = duckdb.__version__
    except Exception:
        return None
    candidate = Path(directory) / f"v{version}" / duckdb_platform() / "excel.duckdb_extension"
    return candidate if candidate.is_file() else None


def load_excel_extension(connection, request: dict) -> bool:
    """LOAD the vendored extension with install/autoload disabled."""
    path = duckdb_extension_path(request)
    if path is None:
        return False
    connection.execute("SET autoinstall_known_extensions=false")
    connection.execute("SET autoload_known_extensions=false")
    try:
        connection.execute(f"LOAD {quote_literal(str(path))}")
        return True
    except Exception:
        return False


def oversize_message(estimate, budget, request: dict) -> str:
    cells = f"{estimate:,}" if estimate else "未知数量"
    lines = [
        f"该工作表约 {cells} 个单元格，超过单次转换预算 {budget:,} 个单元格。整表用 Python 解析会跑几分钟，这里直接拒绝而不是让你等到超时。",
        "可选处理方式（按推荐顺序）：",
        "1) 安装随包 duckdb excel 扩展后即可流式读取（2000 万单元格约 11s、内存约 250MB）：设置 ANYONG_DUCKDB_EXTENSION_DIR，或运行 scripts/install-cnb-duckdb-extension.ps1。",
        "2) 把该表另存为 CSV/Parquet 后重新查询，duckdb 直接扫描，内存最低。",
        "3) 按工作表或列拆分，或先用 sheet_query 指定范围取样本。",
        f"4) 提高 ANYONG_SHEET_CELL_BUDGET（当前 {budget:,}），代价是 calamine 内存约为工作表 XML 的 1.5~2 倍。",
        "5) 明确接受慢速回退：设置 ANYONG_SHEET_FALLBACK=openpyxl（2000 万单元格约 105s，仅在有超时余量时使用）。",
    ]
    if duckdb_extension_path(request) is None and str(request.get("extensionDirectory") or "").strip():
        lines.insert(0, "已配置 ANYONG_DUCKDB_EXTENSION_DIR，但没有找到与当前 duckdb 版本/平台匹配的扩展文件。")
    return "\n".join(lines)


def resolve_reader(request: dict, path: str) -> str:
    """Pick a reader tier: calamine fast path, duckdb excel extension, else explicit failure."""
    forced = str(request.get("reader") or "").strip().lower()
    budget = int(request.get("cellBudget") or DEFAULT_CELL_BUDGET)
    estimate = estimate_cells(path)
    over_budget = estimate is not None and estimate > budget

    if forced in ("calamine", "openpyxl", "duckdb-excel", "duckdb_excel"):
        if forced.startswith("duckdb"):
            return "duckdb-excel"
        if forced == "calamine" and not calamine_available():
            raise SheetError("未安装 python-calamine，无法使用 calamine 读取路径")
        return forced
    if over_budget:
        if duckdb_extension_path(request) is not None:
            return "duckdb-excel"
        if str(request.get("fallback") or "").strip().lower() == "openpyxl":
            return "openpyxl"
        raise SheetError(oversize_message(estimate, budget, request))
    return "calamine" if calamine_available() else "openpyxl"


def cache_paths(source: Path, sheet_key: str, cache_dir: Path):
    stats = source.stat()
    key = f"{source.resolve()}|{stats.st_size}|{int(stats.st_mtime)}|{sheet_key}"
    import hashlib

    digest = hashlib.sha256(key.encode("utf-8")).hexdigest()[:24]
    stem = f"{source.stem[:40]}-{digest}"
    return cache_dir / f"{stem}.csv", cache_dir / f"{stem}.parquet"


def csv_value(value) -> str:
    """Keep integral numbers integral so a column stays BIGINT, not DOUBLE.

    calamine returns every numeric cell as float; without this, an ID or count
    column would silently become a floating point column in the SQL layer.
    """
    if isinstance(value, float) and value.is_integer() and abs(value) <= 2 ** 53:
        return str(int(value))
    return cell_text(value)


def write_csv(rows, target: Path) -> int:
    count = 0
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(".csv.part")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            if row is None:
                continue
            values = [csv_value(value) for value in row]
            while values and values[-1] == "":
                values.pop()
            writer.writerow(values)
            count += 1
    temporary.replace(target)
    return count


def export_with_openpyxl(path: str, sheet_name, target: Path):
    workbook = open_workbook(path)
    sheet, names = select_sheet(workbook, sheet_name)
    rows = write_csv(sheet.iter_rows(values_only=True), target)
    return rows, sheet.title, names


def export_with_calamine(path: str, sheet_name, target: Path):
    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_path(path)
    names = list(workbook.sheet_names)
    if sheet_name:
        if sheet_name not in names:
            raise SheetError(f"工作表 {sheet_name} 不存在；可用工作表：{', '.join(names)}")
        sheet = workbook.get_sheet_by_name(sheet_name)
        title = sheet_name
    else:
        sheet = workbook.get_sheet_by_index(0)
        title = names[0] if names else ""
    rows = write_csv(sheet.iter_rows(), target)
    return rows, title, names


def export_sheet_cache(request: dict, cache_dir: Path):
    """Build the CSV cache and remember which reader produced it."""
    source = Path(request["path"])
    sheet_key = str(request.get("sheet") or "")
    reader = resolve_reader(request, str(source))
    csv_path, parquet_path = cache_paths(source, sheet_key, cache_dir)
    if csv_path.is_file():
        return csv_path, parquet_path, None, None, reader
    try:
        if reader == "calamine":
            rows, title, names = export_with_calamine(str(source), request.get("sheet"), csv_path)
        else:
            rows, title, names = export_with_openpyxl(str(source), request.get("sheet"), csv_path)
    except SheetError:
        raise
    except Exception as error:
        raise SheetError(f"读取表格失败：{error}") from error
    return csv_path, parquet_path, rows, {"title": title, "names": names}, reader


def quote_literal(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


FORBIDDEN_KEYWORDS = (
    "attach", "detach", "copy", "export", "import", "install", "load", "pragma",
    "set ", "reset", "create", "insert", "update", "delete", "drop", "alter",
    "replace", "call", "checkpoint", "vacuum", "truncate", "globs", "read_csv",
    "read_parquet", "read_json", "read_xlsx", "read_text", "sniff_csv",
    "parquet_scan", "delta_scan", "iceberg_scan", "postgres_scan", "sqlite_scan",
)


def validate_select_sql(raw) -> str:
    if not isinstance(raw, str) or not raw.strip():
        raise SheetError("sql 不能为空")
    sql = raw.strip()
    while sql.endswith(";"):
        sql = sql[:-1].rstrip()
    if ";" in sql:
        raise SheetError("只允许单条只读查询；请去掉分号后的额外语句")
    lowered = sql.lower()
    if not (lowered.startswith("select ") or lowered.startswith("select\n") or lowered.startswith("with ")):
        raise SheetError("只允许 SELECT 或 WITH 开头的只读查询")
    for keyword in FORBIDDEN_KEYWORDS:
        if keyword in lowered:
            raise SheetError(f"查询包含不允许的关键字：{keyword.strip()}")
    return sql


def aggregate(request: dict) -> dict:
    try:
        import duckdb
    except ImportError as error:  # pragma: no cover - runtime provisioning issue
        raise SheetError(f"Python 运行时缺少 duckdb：{error}") from error
    limit = int(request.get("limit") or 200)
    limit = max(1, min(limit, MAX_SQL_ROWS))
    cache_dir = Path(request.get("cacheDir") or tempfile.gettempdir())
    cache_dir.mkdir(parents=True, exist_ok=True)
    started = time.time()
    source = Path(request["path"])
    sheet_key = str(request.get("sheet") or "")
    csv_path = Path(request.get("csvPath") or cache_paths(source, sheet_key, cache_dir)[0])
    parquet_path = csv_path.with_suffix(".parquet")
    cache_hit = csv_path.is_file() or parquet_path.is_file()
    reader = None
    sheet_title = request.get("sheet") or ""
    sheet_names = None
    rows_written = None
    direct = False
    if not cache_hit:
        if resolve_reader(request, str(source)) == "duckdb-excel":
            # The extension streams the workbook, so no CSV snapshot is needed.
            direct = True
            reader = "duckdb-excel"
            sheet_names = fast_sheet_names(str(source))
            sheet_title = sheet_title or (sheet_names[0] if sheet_names else "")
        else:
            csv_path, parquet_path, rows_written, sheet_info, reader = export_sheet_cache(request, cache_dir)
            if sheet_info:
                sheet_title = sheet_info["title"]
                sheet_names = sheet_info["names"]
    else:
        sheet_names = fast_sheet_names(str(source))
        if sheet_names:
            sheet_title = sheet_title or sheet_names[0]
    conversion_ms = int((time.time() - started) * 1000)

    memory_limit = request.get("memoryLimit") or "512MB"
    threads = int(request.get("threads") or 2)
    connection = duckdb.connect(database=":memory:")
    try:
        connection.execute(f"SET memory_limit={quote_literal(memory_limit)}")
        connection.execute(f"SET threads={threads}")
        connection.execute(f"SET temp_directory={quote_literal(str(cache_dir / 'spill'))}")
        connection.execute("SET preserve_insertion_order=false")
        sql = request.get("sql")
        if direct:
            if not load_excel_extension(connection, request):
                raise SheetError(oversize_message(estimate_cells(str(source)), int(request.get("cellBudget") or DEFAULT_CELL_BUDGET), request))
            connection.execute(
                "CREATE TABLE data AS SELECT * FROM read_xlsx("
                f"{quote_literal(str(source))}, sheet={quote_literal(str(sheet_title))}, header=true)"
            )
            connection.execute(
                f"COPY (SELECT * FROM data) TO {quote_literal(str(parquet_path))} (FORMAT parquet)"
            )
            connection.execute("SET enable_external_access=false")
            columns = [row[0] for row in connection.execute("DESCRIBE data").fetchall()]
            meta_direct = {
                "sheet": sheet_title,
                "sheets": sheet_names,
                "columns": columns,
                "reader": "duckdb-excel",
                "cellsEstimate": estimate_cells(str(source)),
                "cache": {"hit": False, "csv": False, "parquet": parquet_path.is_file(), "rows": None},
                "conversionMs": int((time.time() - started) * 1000),
            }
            if not sql:
                sample = connection.execute(f"SELECT * FROM data LIMIT {limit}").fetchall()
                return {
                    **meta_direct,
                    "rows": [[json_value(value) for value in row] for row in sample],
                    "rowCount": len(sample),
                    "truncated": len(sample) >= limit,
                    "query": None,
                }
            statement = validate_select_sql(sql)
            cursor = connection.execute(statement)
            fetched = cursor.fetchmany(limit + 1)
            rows = [[json_value(value) for value in row] for row in fetched[:limit]]
            return {
                **meta_direct,
                "rows": rows,
                "rowCount": len(rows),
                "truncated": len(fetched) > limit,
                "query": statement,
                "elapsedMs": int((time.time() - started) * 1000),
            }
        # Parquet keeps repeat runs at tens of milliseconds and shrinks the cache
        # (about 4.5x smaller than the CSV snapshot).
        if not parquet_path.is_file():
            connection.execute(
                "COPY (SELECT * FROM read_csv_auto("
                f"{quote_literal(str(csv_path))}, header=true, sample_size=100000)) "
                f"TO {quote_literal(str(parquet_path))} (FORMAT parquet)"
            )
        connection.execute(
            "CREATE TABLE data AS SELECT * FROM read_parquet("
            f"{quote_literal(str(parquet_path))})"
        )
        columns = [row[0] for row in connection.execute("DESCRIBE data").fetchall()]
        meta = {
            "sheet": sheet_title,
            "sheets": sheet_names,
            "columns": columns,
            "reader": reader or "cache",
            "cellsEstimate": estimate_cells(str(source)),
            "cache": {
                "hit": cache_hit,
                "csv": csv_path.is_file(),
                "parquet": parquet_path.is_file(),
                "rows": rows_written,
            },
            "conversionMs": conversion_ms,
        }
        if not sql:
            sample = connection.execute(f"SELECT * FROM data LIMIT {limit}").fetchall()
            return {
                **meta,
                "rows": [[json_value(value) for value in row] for row in sample],
                "rowCount": len(sample),
                "truncated": len(sample) >= limit,
                "query": None,
            }
        statement = validate_select_sql(sql)
        connection.execute("SET enable_external_access=false")
        cursor = connection.execute(statement)
        fetched = cursor.fetchmany(limit + 1)
        rows = [[json_value(value) for value in row] for row in fetched[:limit]]
        return {
            **meta,
            "rows": rows,
            "rowCount": len(rows),
            "truncated": len(fetched) > limit,
            "query": statement,
            "elapsedMs": int((time.time() - started) * 1000),
        }
    except SheetError:
        raise
    except Exception as error:
        raise SheetError(f"查询失败：{error}") from error
    finally:
        connection.close()


OPERATIONS = {
    "profile": profile,
    "query": query,
    "aggregate": aggregate,
}


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="strict")
    except Exception:
        pass
    try:
        request = read_request()
        operation = request.get("operation")
        handler = OPERATIONS.get(operation)
        if handler is None:
            raise SheetError(f"不支持的操作：{operation}")
        payload = handler(request)
        payload["ok"] = True
        print(json.dumps(payload, ensure_ascii=False, default=str))
        return 0
    except SheetError as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 2
    except Exception:  # pragma: no cover - unexpected engine failure
        print(json.dumps({"ok": False, "error": "表格引擎内部错误", "detail": traceback.format_exc(limit=3)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())